#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""探针：《合同生成》技能端到端验证（真实调模型）

跑一次真实对话，验证 contract-generator 技能：
  1. 模型是否自己读了这个技能的 SKILL.md（技能触发）；
  2. 是否 write_todos 建过程清单；
  3. 是否 docx_list_placeholders 列占位符、read_file 读默认甲方；
  4. 是否 docx_fill_template 填充生成（弹审批卡后自动批准）；
  5. 生成 .docx 无 {{}} 残留，最终回复含结果。

    python _probe_contract.py [自定义提问]
"""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path

import requests

BASE = "http://127.0.0.1:8765"

# --- 会话隔离：会话接口要求声明调用方身份（见 server.py 会话隔离设计）---
# 未带身份时：列表返回空、单会话按「不存在」返回 404。测试脚本必须带上。
_IDENT = {"user_phone": '13912345678', "user_name": '系统管理员'}
_Q = "user_phone=13912345678&user_name=%E7%B3%BB%E7%BB%9F%E7%AE%A1%E7%90%86%E5%91%98"

CONTRACT_DIR = Path(__file__).parent / "static" / "contracts"

DEFAULT_PROMPT = (
    "帮我生成一份课程服务合同：乙方是李四，身份证号 110101199001011234，电话 13900000002；"
    "课程是 AI 大模型实战训练营，服务期限 6 个月，总费用 12800 元。"
)


def sse_events(resp):
    for raw in resp.iter_lines(decode_unicode=True):
        if not raw:
            continue
        line = raw.decode("utf-8") if isinstance(raw, bytes) else raw
        if not line.startswith("data: "):
            continue
        try:
            yield json.loads(line[6:])
        except json.JSONDecodeError:
            continue


def as_dict(raw) -> dict:
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, str):
        try:
            parsed = json.loads(raw)
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}
    return {}


def main() -> int:
    prompt = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PROMPT

    # 清理上次产物
    CONTRACT_DIR.mkdir(parents=True, exist_ok=True)
    for f in CONTRACT_DIR.glob("*.docx"):
        f.unlink(missing_ok=True)

    sid = requests.post(f"{BASE}/api/sessions", json={"title": "E2E-合同", **_IDENT}, timeout=20).json()["id"]
    print(f"会话: {sid}")
    print(f"提问: {prompt}\n" + "-" * 70)

    todos_snapshots: list[list[dict]] = []
    tool_calls: list[tuple[str, dict]] = []
    tool_results: list[tuple[str, str]] = []
    blocked: list[str] = []
    approved = 0
    final_text = ""
    errors: list[str] = []

    with requests.post(
        f"{BASE}/api/chat",
        json={"session_id": sid, "content": prompt},
        stream=True,
        timeout=900,
    ) as resp:
        resp.raise_for_status()
        for ev in sse_events(resp):
            kind = ev.get("event")
            if kind == "llm_token":
                final_text += ev.get("content") or ev.get("token") or ""
            elif kind == "todo":
                todos_snapshots.append(ev.get("todos") or [])
                brief = " | ".join(f"{t.get('status')}:{t.get('content')}" for t in todos_snapshots[-1])
                print(f"[todo #{len(todos_snapshots)}] {brief}")
            elif kind == "tool_start":
                name = ev.get("name") or ev.get("tool") or "?"
                args = as_dict(ev.get("args"))
                tool_calls.append((name, args))
                print(f"[tool] {name}  {json.dumps(args, ensure_ascii=False)[:160]}")
            elif kind == "tool_end":
                tool_results.append((str(ev.get("name") or ""), str(ev.get("result") or "")))
            elif kind == "tool_blocked":
                blocked.append(str(ev.get("tool") or ev.get("name") or "?"))
                print(f"[BLOCKED] {blocked[-1]}")
            elif kind == "approval_request":
                approved += 1
                print(f"[approval #{approved}] 自动批准")
                requests.post(
                    f"{BASE}/api/chat/{sid}/approve",
                    params=_IDENT,
                    json={"approved": True, "session_id": sid},
                    timeout=20,
                )
            elif kind == "error":
                errors.append(str(ev))
                print(f"[ERROR] {json.dumps(ev, ensure_ascii=False)[:300]}")
            elif kind == "done":
                pass

    print("-" * 70)
    names = [n for n, _ in tool_calls]
    read_skill = any(
        n == "read_file" and "contract-generator/SKILL.md" in str(a.get("file_path") or a.get("path") or "")
        for n, a in tool_calls
    )
    read_defaults = any(
        n == "read_file" and "defaults.json" in str(a.get("file_path") or a.get("path") or "")
        for n, a in tool_calls
    )
    listed = "docx_list_placeholders" in names
    filled = "docx_fill_template" in names

    # 找生成的 .docx 并回读验证
    outputs = sorted(CONTRACT_DIR.glob("*.docx"), key=lambda p: p.stat().st_mtime)
    out_ok = False
    if outputs:
        out = outputs[-1]
        try:
            from docx import Document
            txt = "\n".join(p.text for p in Document(str(out)).paragraphs)
            residual = re.findall(r"\{\{[^}]+\}\}", txt)
            out_ok = not residual
            print(f"输出文件: {out.name}  ({out.stat().st_size} 字节)，残留占位符: {len(residual)}")
        except Exception as e:
            print(f"回读输出文件失败: {e}")

    checks = [
        ("读了 contract-generator/SKILL.md（技能触发）", read_skill),
        ("调用了 write_todos", "write_todos" in names),
        ("todo 事件 ≥ 1 次", len(todos_snapshots) >= 1),
        ("调用了 docx_list_placeholders", listed),
        ("读了默认甲方 defaults.json", read_defaults),
        ("调用了 docx_fill_template（写）", filled),
        ("生成文件弹了审批卡（设计如此）", approved >= 1),
        ("生成了 .docx 文件", bool(outputs)),
        ("生成文件无 {{}} 残留", out_ok),
        ("最终回复含结果", bool(final_text.strip())),
        ("无 error 事件", not errors),
    ]

    print("\n工具链:", " -> ".join(names) or "(无)")
    print(f"todo 快照: {len(todos_snapshots)} 次 | 审批: {approved} 次 | 被拦: {blocked or '无'}")
    print("\n===== 断言 =====")
    ok = True
    for label, passed in checks:
        print(f"  {'PASS' if passed else 'FAIL'}  {label}")
        ok = ok and passed
    print("\n最终回复（前 500 字）:\n" + final_text.strip()[:500])
    print(f"\n结果: {'全部通过' if ok else '存在失败项'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
