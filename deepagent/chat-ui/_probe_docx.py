#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""探针：Word（.docx）工具端到端验证（真实调模型）

跑一次真实对话，验证 Agent 是否真的能：
  1. 读取 .docx 合同模板（docx_read_text / docx_list_placeholders）；
  2. 列出模板占位符；
  3. 按示例数据填充模板，生成新 .docx（docx_fill_template，弹审批卡后自动批准）；
  4. 生成的文件里占位符被正确替换（无 {{}} 残留）。

    python _probe_docx.py [自定义提问]
"""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path

import requests

BASE = "http://127.0.0.1:8765"
TPL = r"C:\Users\Administrator\Documents\deepagent\crm_files\课程服务合同word模板.docx"
OUT = r"C:\Users\Administrator\Documents\deepagent\crm_files\_e2e_合同_李四.docx"

DEFAULT_PROMPT = (
    "请帮我用合同模板生成一份合同，分三步：\n"
    f"1. 用 docx_read_text 读取合同模板 {TPL} 的内容，再用 docx_list_placeholders 列出它的占位符；\n"
    "2. 用 docx_fill_template 填充以下数据（JSON 对象）：\n"
    "   {\"CompanyName1\":\"北京智学在线教育科技有限公司\",\"SocialCreditCode1\":\"91110108MA01ABCDEF\","
    "\"Name1\":\"张三\",\"TelNumber1\":\"13800000001\",\"CompanyName2\":\"李四\","
    "\"SocialCreditCode2\":\"110101199001011234\",\"TelNumber2\":\"13900000002\","
    "\"CourseName\":\"AI 大模型实战训练营\",\"ServiceTerm\":\"6\",\"Amount\":\"12800\"}\n"
    f"   输出路径为 {OUT}；\n"
    "3. 用一句话告诉我生成结果（文件名即可，不要贴全文）。"
)


def sse_events(resp):
    for raw in resp.iter_lines(decode_unicode=True):
        if not raw:
            continue
        line = raw.decode("utf-8") if isinstance(raw, bytes) else raw
        if not line.startswith("data: "):
            continue
        try:
            yield json.loads(line[6:])
        except json.JSONDecodeError:
            continue


def as_dict(raw) -> dict:
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, str):
        try:
            parsed = json.loads(raw)
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}
    return {}


def main() -> int:
    prompt = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PROMPT

    # 清理上次输出，避免 docx_fill_template 因「文件已存在」拒绝
    Path(OUT).unlink(missing_ok=True)

    sid = requests.post(f"{BASE}/api/sessions", json={"title": "E2E-Word"}, timeout=20).json()["id"]
    print(f"会话: {sid}")
    print(f"提问: {prompt}\n" + "-" * 70)

    tool_calls: list[tuple[str, dict]] = []
    tool_results: list[tuple[str, str]] = []
    blocked: list[str] = []
    approved = 0
    final_text = ""
    errors: list[str] = []

    with requests.post(
        f"{BASE}/api/chat",
        json={"session_id": sid, "content": prompt},
        stream=True,
        timeout=900,
    ) as resp:
        resp.raise_for_status()
        for ev in sse_events(resp):
            kind = ev.get("event")
            if kind == "llm_token":
                final_text += ev.get("content") or ev.get("token") or ""
            elif kind == "tool_start":
                name = ev.get("name") or ev.get("tool") or "?"
                args = as_dict(ev.get("args"))
                tool_calls.append((name, args))
                print(f"[tool] {name}  {json.dumps(args, ensure_ascii=False)[:200]}")
            elif kind == "tool_end":
                tool_results.append((str(ev.get("name") or ""), str(ev.get("result") or "")))
            elif kind == "tool_blocked":
                blocked.append(str(ev.get("tool") or ev.get("name") or "?"))
                print(f"[BLOCKED] {blocked[-1]}")
            elif kind == "approval_request":
                approved += 1
                print(f"[approval #{approved}] 自动批准")
                requests.post(
                    f"{BASE}/api/chat/{sid}/approve",
                    json={"approved": True, "session_id": sid},
                    timeout=20,
                )
            elif kind == "error":
                errors.append(str(ev))
                print(f"[ERROR] {json.dumps(ev, ensure_ascii=False)[:300]}")
            elif kind == "done":
                pass

    print("-" * 70)
    names = [n for n, _ in tool_calls]
    read_called = any(n in ("docx_read_text", "docx_list_placeholders") for n in names)
    listed = "docx_list_placeholders" in names
    filled = "docx_fill_template" in names

    # 回读生成的文件验证无残留占位符
    out_ok = False
    residual = None
    if Path(OUT).is_file():
        try:
            from docx import Document
            doc = Document(OUT)
            txt = "\n".join(p.text for p in doc.paragraphs)
            residual = re.findall(r"\{\{[^}]+\}\}", txt)
            out_ok = not residual
            print(f"输出文件: {OUT}  ({Path(OUT).stat().st_size} 字节)，残留占位符: {len(residual)}")
        except Exception as e:
            print(f"回读输出文件失败: {e}")

    checks = [
        ("调用了 Word 读/列占位符工具", read_called),
        ("调用了 docx_list_placeholders", listed),
        ("调用了 docx_fill_template（写）", filled),
        ("生成文件弹了审批卡（设计如此）", approved >= 1),
        ("生成了新 .docx 文件", Path(OUT).is_file()),
        ("生成文件无 {{}} 残留", out_ok),
        ("最终回复提到了生成结果", bool(final_text.strip())),
        ("无 error 事件", not errors),
    ]

    print("\n工具链:", " -> ".join(names) or "(无)")
    print(f"审批: {approved} 次 | 被拦: {blocked or '无'}")
    print("\n===== 断言 =====")
    ok = True
    for label, passed in checks:
        print(f"  {'PASS' if passed else 'FAIL'}  {label}")
        ok = ok and passed
    print("\n最终回复（前 500 字）:\n" + final_text.strip()[:500])
    print(f"\n结果: {'全部通过' if ok else '存在失败项'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
