"""Word 文档（.docx）读写工具集。

把 python-docx 封装成 Agent 可直接调用的工具，支撑「读合同模板 → 填占位符 →
生成新合同」这类业务。三个工具：

- ``docx_read_text``          —— 读 .docx 全文纯文本（含表格），理解内容与结构；
- ``docx_list_placeholders``  —— 列出模板里的 ``{{占位符}}`` 及出现次数；
- ``docx_fill_template``      —— 按占位符映射填充模板，另存为新 .docx。

设计要点
--------
- 路径解析与 ``kb_tools._resolve`` 一致：接受绝对路径、虚拟路径（``/chat-ui/...``）、
  普通相对路径三种写法。
- 文件由工具自己写盘（python-docx ``Document.save()``），不经过框架 ``write_file``，
  因此不受「新建文件硬拦截」影响；但 ``docx_fill_template`` 在控制面板默认设为
  「人工审批」档，生成文件前会弹审批卡。
- 占位符替换做了跨 run 兜底：优先 run 级替换（保留原格式），若占位符被 Word
  拆进多个 run，则退化为段落级重建（以首 run 的格式承载全文）。
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from langchain_core.tools import tool

# deepagents 项目根（chat-ui 的上一级）
PROJECT_ROOT: Path = Path(__file__).resolve().parent.parent

# 占位符形如 {{Name}}，允许 key 两侧有空格：{{ Name }}
_PH_RE = re.compile(r"\{\{\s*([^{}\s][^{}]*?)\s*\}\}")


def _resolve(path: str) -> Path:
    """把模型给的路径解析成真实路径（与 kb_tools._resolve 一致）。

    1. 真正的绝对路径（``C:/a/b.docx``）→ 原样使用；
    2. 虚拟路径（``/chat-ui/AGENTS.md``）→ 相对 deepagents 项目根解析；
    3. 普通相对路径（``chat-ui/AGENTS.md``）→ 按项目根解析。
    """
    raw = (path or "").strip().strip('"').strip("'")
    if not raw:
        raise ValueError("路径不能为空")
    p = Path(raw)
    if p.is_absolute():
        return p
    if p.anchor:
        return (PROJECT_ROOT / str(p).lstrip("/\\")).resolve()
    return (PROJECT_ROOT / p).resolve()


def _iter_paragraphs(doc):
    """遍历正文段落 + 所有表格单元格内的段落（合并单元格去重）。"""
    for p in doc.paragraphs:
        yield p
    seen = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                if id(tc) in seen:
                    continue
                seen.add(id(tc))
                for p in cell.paragraphs:
                    yield p


def _paragraph_texts(doc) -> list[str]:
    """正文 + 表格的全部段落文本。"""
    return [p.text for p in _iter_paragraphs(doc)]


def _scan_placeholders(texts) -> dict[str, int]:
    """统计每段文本里出现的 {{占位符}} 次数。"""
    counts: dict[str, int] = {}
    for t in texts:
        for m in _PH_RE.finditer(t):
            key = m.group(1)
            counts[key] = counts.get(key, 0) + 1
    return counts


def _fill_paragraph(paragraph, mapping: dict[str, str]) -> None:
    """替换单个段落里的 ``{{key}}``（保留格式；跨 run 拆分时退化为段落级重建）。"""
    if "{{" not in paragraph.text:
        return

    # 1) run 级替换：占位符完整落在某个 run 内时，直接改 run.text，保留其原有格式。
    for run in paragraph.runs:
        t = run.text
        if "{{" not in t:
            continue
        for key, val in mapping.items():
            t = t.replace("{{" + key + "}}", val)
            t = t.replace("{{ " + key + " }}", val)
        if t != run.text:
            run.text = t

    # 2) 兜底：占位符被 Word 拆进多个 run，run 级替换无能为力 → 段落级重建。
    if "{{" not in paragraph.text:
        return
    new_text = paragraph.text
    for key, val in mapping.items():
        new_text = new_text.replace("{{" + key + "}}", val)
        new_text = new_text.replace("{{ " + key + " }}", val)
    if new_text == paragraph.text:
        return  # 仍然没替换成功，保留原样
    runs = paragraph.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""


@tool
def docx_read_text(path: str) -> str:
    """读取 .docx 文件的全文纯文本（含表格内容），用于理解合同 / 文档写了什么、有哪些条款。

    适用场景：需要知道某个 Word 模板或文档的内容与结构时。
    Args:
        path: .docx 文件路径。支持绝对路径（C:/...）、虚拟路径（/chat-ui/...）、相对路径。
    """
    try:
        fp = _resolve(path)
        if fp.suffix.lower() != ".docx":
            return f"错误：{fp} 不是 .docx 文件（后缀为 {fp.suffix}）"
        if not fp.is_file():
            return f"错误：文件不存在 {fp}"
        doc = Document(str(fp))
        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)
        for ti, table in enumerate(doc.tables):
            parts.append(f"[表格 {ti + 1}]")
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
        return text if text.strip() else "（该文档没有可提取的文本内容）"
    except Exception as e:
        return f"读取 docx 失败: {e}"


@tool
def docx_list_placeholders(path: str) -> str:
    """列出 .docx 模板里的所有 {{占位符}} 及其出现次数。用于填写合同前确认需要填哪些字段。

    Args:
        path: .docx 模板文件路径。
    """
    try:
        fp = _resolve(path)
        if fp.suffix.lower() != ".docx":
            return f"错误：{fp} 不是 .docx 文件（后缀为 {fp.suffix}）"
        if not fp.is_file():
            return f"错误：文件不存在 {fp}"
        doc = Document(str(fp))
        counts = _scan_placeholders(_paragraph_texts(doc))
        if not counts:
            return "该文档中没有找到 {{占位符}}（形如 {{Name}}）。"
        lines = [f"共 {len(counts)} 个占位符："]
        for k, v in counts.items():
            lines.append(f"- {{{{ {k} }}}}  ×{v}")
        return "\n".join(lines)
    except Exception as e:
        return f"列出占位符失败: {e}"


@tool
def docx_fill_template(template_path: str, replacements: str, output_path: str) -> str:
    """按占位符映射填充 .docx 模板，生成一份新的 Word 文件。

    模板里的占位符形如 {{Name}}（例如 {{CompanyName1}}、{{Amount}}）。
    replacements 是 JSON 对象字符串，key 是占位符名（**不带** {{}}），value 是填入文字，例如：
    '{"CompanyName1":"北京某某科技有限公司","Amount":"12800","CourseName":"AI 实战训练营"}'

    生成结果写入 output_path；若该文件已存在会拒绝覆盖，请换一个文件名。
    output_path 支持 ``{timestamp}`` / ``{ts}`` 占位符，会被替换为当前时间
    （格式 YYYYMMDD_HHMMSS），用于生成带时间戳的唯一文件名，无需自行计算时间。
    Args:
        template_path: 模板 .docx 路径。
        replacements: JSON 对象字符串，占位符名 → 填写值。
        output_path: 输出 .docx 的路径（新文件，须以 .docx 结尾）。
    """
    try:
        # 时间戳占位符：由工具生成当前时间，避免模型自行计算时间出错
        _ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_path.replace("{timestamp}", _ts).replace("{ts}", _ts)

        tpl = _resolve(template_path)
        if tpl.suffix.lower() != ".docx":
            return f"错误：模板 {tpl} 不是 .docx 文件"
        if not tpl.is_file():
            return f"错误：模板文件不存在 {tpl}"

        out = _resolve(output_path)
        if out.suffix.lower() != ".docx":
            return f"错误：输出文件需以 .docx 结尾，得到 {out}"
        if out.exists():
            return f"错误：输出文件已存在 {out}，为避免覆盖请换一个文件名。"

        # 解析替换映射（容错：模型可能直接传 dict）
        if isinstance(replacements, dict):
            mapping = {str(k): str(v) for k, v in replacements.items()}
        else:
            try:
                data = json.loads(replacements)
            except json.JSONDecodeError as e:
                return f"错误：replacements 不是合法 JSON：{e}"
            if not isinstance(data, dict):
                return "错误：replacements 必须是 JSON 对象（占位符名 → 值）。"
            mapping = {str(k): str(v) for k, v in data.items()}
        if not mapping:
            return "错误：replacements 为空，没有可填充的内容。"

        doc = Document(str(tpl))
        present = _scan_placeholders(_paragraph_texts(doc))

        for p in _iter_paragraphs(doc):
            _fill_paragraph(p, mapping)

        remaining = _scan_placeholders(_paragraph_texts(doc))

        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))

        filled = {k: present[k] for k in present if k not in remaining}
        missing_keys = [k for k in mapping if k not in present]
        lines = [f"已生成：{out}"]
        if filled:
            lines.append("已填充：" + ", ".join(f"{k}(×{v})" for k, v in filled.items()))
        if remaining:
            lines.append("⚠️ 仍残留未填充的占位符：" + ", ".join(f"{k}(×{v})" for k, v in remaining.items()))
        if missing_keys:
            lines.append("⚠️ 以下键在模板中不存在，已忽略：" + ", ".join(missing_keys))
        return "\n".join(lines)
    except Exception as e:
        return f"填充模板失败: {e}"


# 工具清单与名称集合（供 server.py 汇总到 base_tools）
DOCX_TOOLS = [docx_read_text, docx_list_placeholders, docx_fill_template]
DOCX_TOOL_NAMES = {t.name for t in DOCX_TOOLS}
